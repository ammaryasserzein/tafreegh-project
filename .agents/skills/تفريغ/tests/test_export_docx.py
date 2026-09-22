import unittest
from pathlib import Path
import sys

# Add scripts directory to sys.path
scripts_dir = Path(__file__).parent.parent / "scripts"
sys.path.insert(0, str(scripts_dir))

from export_docx import get_onedrive_folder

class TestGetOneDriveFolder(unittest.TestCase):
    def setUp(self):
        self.base_path = Path("C:/Users/L/OneDrive/1. دوري")

    def test_known_keywords_mapping(self):
        test_cases = [
            ("اسماء", "1. الاسماء الحسنى"),
            ("أسماء", "1. الاسماء الحسنى"),
            ("رياض", "1. رياض الصالحين"),
            ("سيرة", "1. سيرة (الرحيق المختوم)"),
            ("معاملات", "1. معاملات"),
            ("دليل", "1. معاملات"),
            ("توحيد", "2. التوحيد"),
            ("لب", "2. لب الاصول"),
            ("لب الاصول", "2. لب الاصول"),
            ("أصول", "2. لب الاصول"),
            ("اصول", "2. لب الاصول"),
            ("الاصول", "2. لب الاصول"),
            ("الأصول", "2. لب الاصول"),
            ("ديوان", "ديوان الشافعي"),
            ("زاد", "زاد المعاد"),
            ("روضة", "شرح مختصر الروضة"),
            ("مختصر الروضة", "شرح مختصر الروضة"),
            ("صيد", "صيد الخاطر"),
        ]
        for keyword, expected_folder in test_cases:
            with self.subTest(keyword=keyword):
                result = get_onedrive_folder(keyword, self.base_path)
                self.assertEqual(result, self.base_path / expected_folder)

    def test_fallback_creates_new_folder(self):
        keyword = "تفسير"
        result = get_onedrive_folder(keyword, self.base_path)
        self.assertEqual(result, self.base_path / "تفسير")

from export_docx import parse_metadata

class TestParseMetadata(unittest.TestCase):
    def test_extract_metadata_from_standard_header(self):
        sample = """بسم الله الرحمن الرحيم
اسم المادة: سيرة (الرحيق المختوم)
2026-09-10

والحمد لله، والصلاة والسلام على رسول الله..."""
        meta = parse_metadata(sample)
        self.assertEqual(meta["date"], "2026-09-10")
        self.assertEqual(meta["keyword"], "سيرة")
        self.assertEqual(meta["subject_name"], "سيرة (الرحيق المختوم)")

    def test_extract_metadata_from_alt_header(self):
        sample = """بسم الله الرحمن الرحيم
المادة: زاد المعاد (فصل في حجة أبي بكر الصديق رضي الله عنه)
2026-09-18

بسم الله والحمد لله والصلاة والسلام على رسول الله..."""
        meta = parse_metadata(sample)
        self.assertEqual(meta["date"], "2026-09-18")
        self.assertEqual(meta["keyword"], "زاد")
        self.assertEqual(meta["subject_name"], "زاد المعاد")

    def test_extract_metadata_from_filename_token(self):
        sample = "2026-09-12 معاملات.mp3 دليل المصدر والحمد لله والصلاة والسلام..."
        meta = parse_metadata(sample)
        self.assertEqual(meta["date"], "2026-09-12")
        self.assertEqual(meta["keyword"], "معاملات")

    def test_canonical_subject_name_for_tawheed(self):
        sample = """بسم الله الرحمن الرحيم
المادة: توحيد
2026-09-18
"""
        meta = parse_metadata(sample)
        self.assertEqual(meta["subject_name"], "فتح الباري (كتاب التوحيد)")

import tempfile
import docx
from export_docx import create_docx
from docx.oxml.ns import qn

from docx.enum.text import WD_ALIGN_PARAGRAPH

# Ensure UTF-8 output on Windows
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

class TestCreateDocx(unittest.TestCase):
    def test_create_docx_produces_valid_formatted_file(self):
        sample_md = """بسم الله الرحمن الرحيم
المادة: زاد المعاد (فصل في حجة أبي بكر الصديق رضي الله عنه)
2026-09-10

الحمد لله، والصلاة والسلام على رسول الله.

**(فَفِي مُسْنَدِ أَحْمَدَ وَالشَّافِعِيِّ رَضِيَ اللَّهُ عَنْهُمَا أَنَّهُمْ)**

الكفار.

طالب: في الحجة سيدنا.
"""
        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = Path(tmpdir) / "test_out.docx"
            create_docx(sample_md, out_file)
            
            self.assertTrue(out_file.exists())
            self.assertGreater(out_file.stat().st_size, 1000) # Valid docx file
            
            doc = docx.Document(str(out_file))
            self.assertGreaterEqual(len(doc.paragraphs), 5)
            
            # Check 0.5-inch margins from Word base.docx
            section = doc.sections[0]
            self.assertAlmostEqual(section.top_margin.inches, 0.5, places=2)
            self.assertAlmostEqual(section.bottom_margin.inches, 0.5, places=2)
            self.assertAlmostEqual(section.left_margin.inches, 0.5, places=2)
            self.assertAlmostEqual(section.right_margin.inches, 0.5, places=2)
            
            # Check centered 3-line header
            self.assertEqual(doc.paragraphs[0].text.strip(), "بسم الله الرحمن الرحيم")
            self.assertEqual(doc.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
            
            # Line 2 MUST be clean subject name only: 'المادة: زاد المعاد'
            self.assertEqual(doc.paragraphs[1].text.strip(), "المادة: زاد المعاد")
            self.assertEqual(doc.paragraphs[1].alignment, WD_ALIGN_PARAGRAPH.CENTER)
            
            self.assertEqual(doc.paragraphs[2].text.strip(), "2026-09-10")
            self.assertEqual(doc.paragraphs[2].alignment, WD_ALIGN_PARAGRAPH.CENTER)
            
            # Check bold matn run has both w:b and w:bCs
            matn_found = False
            for p in doc.paragraphs:
                for r in p.runs:
                    self.assertIsNone(r.font.color.rgb)
                    if "فَفِي مُسْنَدِ أَحْمَدَ" in r.text:
                        matn_found = True
                        self.assertTrue(r.bold)
                        rPr = r._r.get_or_add_rPr()
                        self.assertIsNotNone(rPr.find(qn('w:bCs')), "w:bCs missing on Arabic bold run")
            self.assertTrue(matn_found, "Matn text was not found with bold formatting")

    def test_student_label_is_never_bold(self):
        sample_md = """بسم الله الرحمن الرحيم
المادة: زاد المعاد
2026-09-18

الحمد لله.

طالب: صوت غير مسموع.
"""
        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = Path(tmpdir) / "test_student.docx"
            create_docx(sample_md, out_file)
            doc = docx.Document(str(out_file))
            
            student_p = [p for p in doc.paragraphs if "طالب:" in p.text]
            self.assertTrue(len(student_p) > 0)
            for r in student_p[0].runs:
                self.assertFalse(r.bold, f"Student run '{r.text}' should NOT be bold")

    def test_header_with_blank_lines_does_not_leak_into_body(self):
        sample_md = """بسم الله الرحمن الرحيم

المادة: فتح الباري (كتاب التوحيد)

2026-09-18

الحمد لله، والصلاة والسلام على رسول الله.
"""
        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = Path(tmpdir) / "test_header.docx"
            create_docx(sample_md, out_file)
            doc = docx.Document(str(out_file))
            
            # Exactly 3 header paragraphs + 1 body paragraph = 4 paragraphs
            self.assertEqual(len(doc.paragraphs), 4)
            self.assertEqual(doc.paragraphs[0].text.strip(), "بسم الله الرحمن الرحيم")
            self.assertEqual(doc.paragraphs[1].text.strip(), "المادة: فتح الباري (كتاب التوحيد)")
            self.assertEqual(doc.paragraphs[2].text.strip(), "2026-09-18")
            self.assertEqual(doc.paragraphs[3].text.strip(), "الحمد لله، والصلاة والسلام على رسول الله.")

from export_docx import export_dual_copies

class TestExportDualCopies(unittest.TestCase):
    def test_export_dual_copies_creates_both_files(self):
        sample_md = """بسم الله الرحمن الرحيم
اسم المادة: سيرة (الرحيق المختوم)
2026-09-10

الحمد لله، والصلاة والسلام على رسول الله.
"""
        with tempfile.TemporaryDirectory() as tmp_onedrive, tempfile.TemporaryDirectory() as tmp_project:
            res = export_dual_copies(
                markdown_text=sample_md,
                base_onedrive=Path(tmp_onedrive),
                project_root=Path(tmp_project)
            )
            
            onedrive_file = res["onedrive_file"]
            project_file = res["project_file"]
            
            # 1. Check OneDrive file: named exactly {date}.docx inside the mapped folder
            expected_od_folder = Path(tmp_onedrive) / "1. سيرة (الرحيق المختوم)"
            self.assertEqual(onedrive_file.parent, expected_od_folder)
            self.assertEqual(onedrive_file.name, "2026-09-10.docx")
            self.assertTrue(onedrive_file.exists())
            
            # 2. Check Project file: saved inside 03_مخرجات_الوورد with descriptive name
            expected_proj_folder = Path(tmp_project) / "03_مخرجات_الوورد"
            self.assertEqual(project_file.parent, expected_proj_folder)
            self.assertTrue(project_file.name.startswith("2026-09-10_سيرة"))
            self.assertTrue(project_file.exists())

if __name__ == "__main__":
    unittest.main()



