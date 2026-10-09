"""Every source line must become its own Word paragraph (^p), never a manual line break (^l).

Real paragraphs let the user jump between lines with Ctrl+Down in Word.
"""
import sys
import tempfile
import unittest
from pathlib import Path

scripts_dir = Path(__file__).parent.parent / "scripts"
sys.path.insert(0, str(scripts_dir))

import docx  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from export_docx import create_docx  # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

SAMPLE_MD = """بسم الله الرحمن الرحيم
المادة: الأسماء الحسنى
2026-09-26

السطر الأول من الفقرة.
**(نص المتن في سطر مستقل)**
طالب: صوت غير مسموع.

فقرة مستقلة بعد سطر فارغ.
"""


class TestParagraphBreaks(unittest.TestCase):
    def _build(self) -> docx.document.Document:
        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = Path(tmpdir) / "breaks.docx"
            create_docx(SAMPLE_MD, out_file)
            return docx.Document(str(out_file))

    def test_no_manual_line_breaks_in_body(self) -> None:
        doc = self._build()
        breaks = doc.element.body.findall(".//" + qn("w:br"))
        self.assertEqual(len(breaks), 0, "Body must not contain w:br (^l) line breaks")

    def test_each_line_is_its_own_paragraph(self) -> None:
        doc = self._build()
        body = [p.text.strip() for p in doc.paragraphs[3:] if p.text.strip()]
        self.assertEqual(
            body,
            [
                "السطر الأول من الفقرة.",
                "(نص المتن في سطر مستقل)",
                "طالب: صوت غير مسموع.",
                "فقرة مستقلة بعد سطر فارغ.",
            ],
        )

    def test_no_empty_paragraphs_injected_around_matn(self) -> None:
        doc = self._build()
        # All paragraphs after the 3 headers must be non-empty (no ^p^p double gaps)
        body_all = [p.text.strip() for p in doc.paragraphs[3:]]
        self.assertEqual(
            body_all,
            [
                "السطر الأول من الفقرة.",
                "(نص المتن في سطر مستقل)",
                "طالب: صوت غير مسموع.",
                "فقرة مستقلة بعد سطر فارغ.",
            ],
            "No empty paragraphs should be injected around standalone matn",
        )
        self.assertEqual(len(doc.paragraphs), 7, "Total paragraphs must equal 3 headers + 4 body lines without empty gap paragraphs")

    def test_split_lines_keep_rtl_and_bold(self) -> None:
        doc = self._build()
        matn_p = next(p for p in doc.paragraphs if "نص المتن" in p.text)
        self.assertTrue(all(r.bold for r in matn_p.runs if r.text.strip()))
        for p in doc.paragraphs[3:]:
            p_rpr = p._p.pPr.find(qn("w:rPr"))
            self.assertIsNotNone(p_rpr.find(qn("w:rtl")), f"RTL missing on: {p.text}")


if __name__ == "__main__":
    unittest.main()
