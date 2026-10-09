import unittest
from pathlib import Path
import sys
import tempfile
import docx
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Ensure UTF-8 output on Windows
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# Add scripts directory to sys.path
scripts_dir = Path(__file__).parent.parent / "scripts"
sys.path.insert(0, str(scripts_dir))

from export_docx import (
    get_onedrive_folder,
    parse_metadata,
    is_header_line,
    standardize_transcript_header,
    create_docx,
    export_documents,
    export_dual_copies,
    archive_processed_inputs,
    is_matching_stem,
    normalize_stem,
    _extract_candidate_query,
    _format_segment_output,
    CANONICAL_SUBJECT_NAMES,
)
from lecture_identity import (
    LectureIdentity,
    load_canonical_routing,
)
from matn_matcher import MatnMatchResult, MatchMode

class TestGetOneDriveFolder(unittest.TestCase):
    def setUp(self):
        self.base_path = Path("C:/Users/L/OneDrive/1. دوري")

    def test_known_keywords_mapping(self):
        test_cases = [
            ("اسماء", "1. الاسماء الحسنى"),
            ("أسماء", "1. الاسماء الحسنى"),
            ("رياض", "1. رياض الصالحين"),
            ("سيرة", "1. سيرة (الرحيق المختوم)"),
            ("معاملات", "1. معاملات"),
            ("دليل", "1. معاملات"),
            ("توحيد", "2. التوحيد"),
            ("لب", "2. لب الاصول"),
            ("لب الاصول", "2. لب الاصول"),
            ("أصول", "2. لب الاصول"),
            ("اصول", "2. لب الاصول"),
            ("الاصول", "2. لب الاصول"),
            ("الأصول", "2. لب الاصول"),
            ("ديوان", "ديوان الشافعي"),
            ("زاد", "زاد المعاد"),
            ("روضة", "شرح مختصر الروضة"),
            ("مختصر الروضة", "شرح مختصر الروضة"),
            ("صيد", "صيد الخاطر"),
            ("منطق", "المنطق"),
            ("المنطق", "المنطق"),
            ("بيوع", "فقه البيوع (سعد الخثلان)"),
            ("البيوع", "فقه البيوع (سعد الخثلان)"),
            ("فقه البيوع", "فقه البيوع (سعد الخثلان)"),
            ("فقه_البيوع", "فقه البيوع (سعد الخثلان)"),
            ("خثلان", "فقه البيوع (سعد الخثلان)"),
        ]
        for keyword, expected_folder in test_cases:
            with self.subTest(keyword=keyword):
                result = get_onedrive_folder(keyword, self.base_path)
                self.assertEqual(result, self.base_path / expected_folder)

    def test_fallback_creates_new_folder(self):
        keyword = "تفسير"
        result = get_onedrive_folder(keyword, self.base_path)
        self.assertEqual(result, self.base_path / "تفسير")


class TestParseMetadata(unittest.TestCase):
    def test_extract_metadata_from_standard_header(self):
        sample = """بسم الله الرحمن الرحيم
اسم المادة: سيرة (الرحيق المختوم)
2026-09-10

والحمد لله، والصلاة والسلام على رسول الله..."""
        meta = parse_metadata(sample)
        self.assertEqual(meta["date"], "2026-09-10")
        self.assertEqual(meta["keyword"], "سيرة")
        self.assertEqual(meta["subject_name"], "سيرة (الرحيق المختوم)")

    def test_extract_metadata_from_alt_header(self):
        sample = """بسم الله الرحمن الرحيم
المادة: زاد المعاد (فصل في حجة أبي بكر الصديق رضي الله عنه)
2026-09-18

بسم الله والحمد لله والصلاة والسلام على رسول الله..."""
        meta = parse_metadata(sample)
        self.assertEqual(meta["date"], "2026-09-18")
        self.assertEqual(meta["keyword"], "زاد")
        self.assertEqual(meta["subject_name"], "زاد المعاد")

    def test_extract_metadata_from_filename_token(self):
        sample = "2026-09-12 معاملات.mp3 دليل المصدر والحمد لله والصلاة والسلام..."
        meta = parse_metadata(sample)
        self.assertEqual(meta["date"], "2026-09-12")
        self.assertEqual(meta["keyword"], "معاملات")

    def test_canonical_subject_name_for_tawheed(self):
        sample = """بسم الله الرحمن الرحيم
المادة: توحيد
2026-09-18
"""
        meta = parse_metadata(sample)
        self.assertEqual(meta["subject_name"], "فتح الباري (كتاب التوحيد)")

    def test_notebooklm_arabic_source_guide_mantiq(self):
        sample = """2026-09-26 منطق.mp3
دليل المصدر
بسم الله الرحمن الرحيم الحمد لله رب العالمين..."""
        meta = parse_metadata(sample)
        self.assertEqual(meta["date"], "2026-09-26")
        self.assertEqual(meta["keyword"], "منطق")
        self.assertEqual(meta["subject_name"], "المنطق (السلم المنورق)")
        self.assertEqual(meta["audio_file"], "2026-09-26 منطق.mp3")

    def test_notebooklm_english_source_guide_seerah(self):
        sample = """2026-10-01 سيرة.mp3

Source guide

الحمد لله رب العالمين..."""
        meta = parse_metadata(sample)
        self.assertEqual(meta["date"], "2026-10-01")
        self.assertEqual(meta["keyword"], "سيرة")
        self.assertEqual(meta["subject_name"], "سيرة (الرحيق المختوم)")
        self.assertEqual(meta["audio_file"], "2026-10-01 سيرة.mp3")

    def test_notebooklm_sources_tag(self):
        sample = """المصادر
2026-10-03 معاملات.mp3
دليل المصدر
الحمد لله..."""
        meta = parse_metadata(sample)
        self.assertEqual(meta["date"], "2026-10-03")
        self.assertEqual(meta["keyword"], "معاملات")
        self.assertEqual(meta["subject_name"], "دليل الطالب (كتاب البيع)")
        self.assertEqual(meta["audio_file"], "2026-10-03 معاملات.mp3")

    def test_numbered_stem_metadata(self):
        sample = "17-فقه_البيوع_17.mp3\nSource guide\nالحمد لله..."
        meta = parse_metadata(sample)
        self.assertIsNone(meta["date"])
        self.assertEqual(meta["lecture_number"], "17")
        self.assertIn(meta["keyword"], ("فقه_البيوع", "فقه البيوع"))
        self.assertEqual(meta["subject_name"], "فقه البيوع (سعد الخثلان)")
        self.assertEqual(meta["audio_file"], "17-فقه_البيوع_17.mp3")

    def test_numbered_stem_with_explicit_headers(self):
        sample = """بسم الله الرحمن الرحيم
المادة: فقه البيوع (سعد الخثلان)
المحاضرة: 17

الحمد لله رب العالمين..."""
        meta = parse_metadata(sample)
        self.assertIsNone(meta["date"])
        self.assertEqual(meta["lecture_number"], "17")
        self.assertEqual(meta["subject_name"], "فقه البيوع (سعد الخثلان)")


class TestHeaderFiltering(unittest.TestCase):
    def test_is_header_line_notebooklm_tags(self):
        self.assertTrue(is_header_line("المصادر"))
        self.assertTrue(is_header_line("مصادر"))
        self.assertTrue(is_header_line("دليل المصدر"))
        self.assertTrue(is_header_line("دليل مصادر"))
        self.assertTrue(is_header_line("Source guide"))
        self.assertTrue(is_header_line("Source Guide"))
        self.assertTrue(is_header_line("Sources"))
        self.assertTrue(is_header_line("2026-10-01 سيرة.mp3"))
        self.assertTrue(is_header_line("2026-09-26 منطق.mp3"))
        self.assertTrue(is_header_line("2026-10-03 معاملات.m4a"))
        self.assertFalse(is_header_line("الحمد لله رب العالمين"))
        self.assertFalse(is_header_line("طالب: صوت غير مسموع."))

    def test_standardize_transcript_header_excludes_notebooklm_tags(self):
        raw = """2026-10-01 سيرة.mp3

Source guide

الحمد لله رب العالمين، والصلاة والسلام على رسول الله."""
        clean = standardize_transcript_header(raw)
        self.assertNotIn("Source guide", clean)
        self.assertNotIn(".mp3", clean)
        self.assertTrue(clean.startswith("بسم الله الرحمن الرحيم\nالمادة: سيرة (الرحيق المختوم)\n2026-10-01"))
        self.assertIn("الحمد لله رب العالمين", clean)

    def test_is_header_line_numbered_headers(self):
        self.assertTrue(is_header_line("المحاضرة: 17"))
        self.assertTrue(is_header_line("المحاضرة: 5"))
        self.assertTrue(is_header_line("17-فقه_البيوع_17.mp3"))
        self.assertTrue(is_header_line("17-فقه_البيوع_17"))

    def test_standardize_transcript_header_numbered_lecture(self):
        raw = """17-فقه_البيوع_17.mp3

Source guide

المادة: فقه البيوع (سعد الخثلان)
المحاضرة: 17

الحمد لله رب العالمين، والصلاة والسلام على رسول الله."""
        clean = standardize_transcript_header(raw)
        self.assertNotIn("Source guide", clean)
        self.assertNotIn(".mp3", clean)
        self.assertTrue(clean.startswith("بسم الله الرحمن الرحيم\nالمادة: فقه البيوع (سعد الخثلان)\nالمحاضرة: 17"))
        self.assertIn("الحمد لله رب العالمين", clean)

    def test_standardize_transcript_header_override_precedence(self):
        # Caller's explicit identity override must take precedence over markdown header
        raw = """بسم الله الرحمن الرحيم
المادة: قديم
2026-01-01

الحمد لله رب العالمين."""
        override = LectureIdentity(lecture_number="5", subject_name="فقه جديد")
        clean = standardize_transcript_header(raw, identity=override)
        self.assertTrue(clean.startswith("بسم الله الرحمن الرحيم\nالمادة: فقه جديد\nالمحاضرة: 5"))
        self.assertNotIn("قديم", clean)


class TestCreateDocx(unittest.TestCase):
    def test_create_docx_override_precedence(self):
        sample_md = """بسم الله الرحمن الرحيم
المادة: قديم
2026-01-01

الحمد لله رب العالمين."""
        with tempfile.TemporaryDirectory() as tmp_dir:
            out_file = Path(tmp_dir) / "output.docx"
            override = LectureIdentity(lecture_number="5", subject_name="فقه جديد")
            create_docx(sample_md, out_file, identity=override)
            doc = docx.Document(str(out_file))
            p_texts = [p.text for p in doc.paragraphs]
            self.assertIn("المادة: فقه جديد", p_texts)
            self.assertIn("المحاضرة: 5", p_texts)
            self.assertNotIn("المادة: قديم", p_texts)

    def test_create_docx_produces_valid_formatted_file(self):
        sample_md = """بسم الله الرحمن الرحيم
المادة: زاد المعاد (فصل في حجة أبي بكر الصديق رضي الله عنه)
2026-09-10

الحمد لله، والصلاة والسلام على رسول الله.

**(فَفِي مُسْنَدِ أَحْمَدَ وَالشَّافِعِيِّ رَضِيَ اللَّهُ عَنْهُمَا أَنَّهُمْ)**

الكفار.

طالب: في الحجة سيدنا.
"""
        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = Path(tmpdir) / "test_out.docx"
            create_docx(sample_md, out_file)
            
            self.assertTrue(out_file.exists())
            self.assertGreater(out_file.stat().st_size, 1000)
            
            doc = docx.Document(str(out_file))
            self.assertGreaterEqual(len(doc.paragraphs), 5)
            
            # Margins
            section = doc.sections[0]
            self.assertAlmostEqual(section.top_margin.inches, 0.5, places=2)
            self.assertAlmostEqual(section.bottom_margin.inches, 0.5, places=2)
            self.assertAlmostEqual(section.left_margin.inches, 0.5, places=2)
            self.assertAlmostEqual(section.right_margin.inches, 0.5, places=2)
            
            # Header
            self.assertEqual(doc.paragraphs[0].text.strip(), "بسم الله الرحمن الرحيم")
            self.assertEqual(doc.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertEqual(doc.paragraphs[1].text.strip(), "المادة: زاد المعاد")
            self.assertEqual(doc.paragraphs[1].alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertEqual(doc.paragraphs[2].text.strip(), "2026-09-10")
            self.assertEqual(doc.paragraphs[2].alignment, WD_ALIGN_PARAGRAPH.CENTER)
            
            # Bold matn
            matn_found = False
            for p in doc.paragraphs:
                for r in p.runs:
                    self.assertIsNone(r.font.color.rgb)
                    if "فَفِي مُسْنَدِ أَحْمَدَ" in r.text:
                        matn_found = True
                        self.assertTrue(r.bold)
                        rPr = r._r.get_or_add_rPr()
                        self.assertIsNotNone(rPr.find(qn('w:bCs')), "w:bCs missing on Arabic bold run")
            self.assertTrue(matn_found, "Matn text was not found with bold formatting")

    def test_student_label_is_never_bold(self):
        sample_md = """بسم الله الرحمن الرحيم
المادة: زاد المعاد
2026-09-18

الحمد لله.

طالب: صوت غير مسموع.
"""
        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = Path(tmpdir) / "test_student.docx"
            create_docx(sample_md, out_file)
            doc = docx.Document(str(out_file))
            
            student_p = [p for p in doc.paragraphs if "طالب:" in p.text]
            self.assertTrue(len(student_p) > 0)
            for r in student_p[0].runs:
                self.assertFalse(r.bold, f"Student run '{r.text}' should NOT be bold")

    def test_header_with_blank_lines_does_not_leak_into_body(self):
        sample_md = """بسم الله الرحمن الرحيم

المادة: فتح الباري (كتاب التوحيد)

2026-09-18

الحمد لله، والصلاة والسلام على رسول الله.
"""
        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = Path(tmpdir) / "test_header.docx"
            create_docx(sample_md, out_file)
            doc = docx.Document(str(out_file))
            
            self.assertEqual(len(doc.paragraphs), 4)
            self.assertEqual(doc.paragraphs[0].text.strip(), "بسم الله الرحمن الرحيم")
            self.assertEqual(doc.paragraphs[1].text.strip(), "المادة: فتح الباري (كتاب التوحيد)")
            self.assertEqual(doc.paragraphs[2].text.strip(), "2026-09-18")
            self.assertEqual(doc.paragraphs[3].text.strip(), "الحمد لله، والصلاة والسلام على رسول الله.")

    def test_create_docx_strips_notebooklm_header_lines_from_body(self):
        sample = """2026-10-01 سيرة.mp3

Source guide

الحمد لله رب العالمين، والصلاة والسلام على رسول الله.
"""
        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = Path(tmpdir) / "test_nb.docx"
            create_docx(sample, out_file)
            doc = docx.Document(str(out_file))
            body_texts = [p.text.strip() for p in doc.paragraphs]
            self.assertEqual(body_texts[0], "بسم الله الرحمن الرحيم")
            self.assertEqual(body_texts[1], "المادة: سيرة (الرحيق المختوم)")
            self.assertEqual(body_texts[2], "2026-10-01")
            self.assertNotIn("Source guide", body_texts)
            self.assertNotIn("2026-10-01 سيرة.mp3", body_texts)
            self.assertEqual(body_texts[3], "الحمد لله رب العالمين، والصلاة والسلام على رسول الله.")


class TestArchiveProcessedInputs(unittest.TestCase):
    def test_archive_moves_matching_raw_and_matn_and_deletes_chunks(self):
        with tempfile.TemporaryDirectory() as tmp_root:
            root = Path(tmp_root)
            matn_dir = root / "01_Matn_Sources"
            raw_dir = root / "02_Raw_Inputs"
            matn_dir.mkdir()
            raw_dir.mkdir()

            # Matching files
            (matn_dir / "2026-10-01__سيرة.md").write_text("matn content", encoding="utf-8")
            (raw_dir / "2026-10-01 سيرة.md").write_text("raw content", encoding="utf-8")
            chunks_dir = raw_dir / "2026-10-01  سيرة.md_مقاطع"
            chunks_dir.mkdir()
            (chunks_dir / "part1.txt").write_text("part 1", encoding="utf-8")

            # Other files
            (matn_dir / "2026-10-03__معاملات.md").write_text("other matn", encoding="utf-8")
            (raw_dir / "2026-10-03 معاملات.md").write_text("other raw", encoding="utf-8")

            res = archive_processed_inputs(LectureIdentity(date="2026-10-01", keyword="سيرة"), project_root=root)

            # Assert moved files exist in processed/
            self.assertTrue((matn_dir / "processed" / "2026-10-01__سيرة.md").exists())
            self.assertFalse((matn_dir / "2026-10-01__سيرة.md").exists())
            self.assertTrue((raw_dir / "processed" / "2026-10-01 سيرة.md").exists())
            self.assertFalse((raw_dir / "2026-10-01 سيرة.md").exists())

            # Assert chunk folder deleted
            self.assertFalse(chunks_dir.exists())

            # Assert other files untouched
            self.assertTrue((matn_dir / "2026-10-03__معاملات.md").exists())
            self.assertTrue((raw_dir / "2026-10-03 معاملات.md").exists())


class TestExportDocuments(unittest.TestCase):
    def test_export_documents_creates_both_files(self):
        sample_md = """بسم الله الرحمن الرحيم
اسم المادة: سيرة (الرحيق المختوم)
2026-09-10

الحمد لله، والصلاة والسلام على رسول الله.
"""
        with tempfile.TemporaryDirectory() as tmp_onedrive, tempfile.TemporaryDirectory() as tmp_project:
            (Path(tmp_onedrive) / "1. سيرة (الرحيق المختوم)").mkdir()
            res = export_documents(
                markdown_text=sample_md,
                base_onedrive=Path(tmp_onedrive),
                project_root=Path(tmp_project)
            )
            
            onedrive_file = res["onedrive_file"]
            project_file = res["project_file"]
            
            expected_od_folder = Path(tmp_onedrive) / "1. سيرة (الرحيق المختوم)"
            self.assertEqual(onedrive_file.parent, expected_od_folder)
            self.assertEqual(onedrive_file.name, "2026-09-10.docx")
            self.assertTrue(onedrive_file.exists())
            
            expected_proj_folder = Path(tmp_project) / "03_AI_Outputs"
            self.assertEqual(project_file.parent, expected_proj_folder)
            self.assertTrue(project_file.name.startswith("2026-09-10_سيرة"))
            self.assertTrue(project_file.exists())


class TestIsMatchingStem(unittest.TestCase):
    def test_normalize_stem_strips_extensions_and_suffixes(self):
        self.assertEqual(
            normalize_stem("2026-09-19_منطق_(السلم_المنورق)_AI_full.md.md"),
            "2026-09-19_منطق_(السلم_المنورق)"
        )
        self.assertEqual(
            normalize_stem("2026-10-01_سيرة_processed.docx"),
            "2026-10-01_سيرة"
        )
        self.assertEqual(
            normalize_stem("2026-10-01  سيرة.md_مقاطع"),
            "2026-10-01  سيرة"
        )
        self.assertEqual(
            normalize_stem("2026-08-29_لب_الأصول_AI.md"),
            "2026-08-29_لب_الأصول"
        )

    def test_is_matching_stem_standard_match(self):
        self.assertTrue(is_matching_stem("2026-10-01_سيرة.md", LectureIdentity(date="2026-10-01", keyword="سيرة")))

    def test_is_matching_stem_double_extension(self):
        self.assertTrue(is_matching_stem("2026-10-01_سيرة.md.md", LectureIdentity(date="2026-10-01", keyword="سيرة")))

    def test_is_matching_stem_suffixes_ai_full_and_processed(self):
        self.assertTrue(is_matching_stem("2026-09-19_منطق_(السلم_المنورق)_AI_full.md", LectureIdentity(date="2026-09-19", keyword="منطق")))
        self.assertTrue(is_matching_stem("2026-10-01_سيرة_processed.docx", LectureIdentity(date="2026-10-01", keyword="سيرة")))
        self.assertTrue(is_matching_stem("2026-08-29_لب_الأصول_AI.md", LectureIdentity(date="2026-08-29", keyword="لب")))

    def test_is_matching_stem_chunk_folder_suffix(self):
        self.assertTrue(is_matching_stem("2026-10-01  سيرة.md_مقاطع", LectureIdentity(date="2026-10-01", keyword="سيرة")))

    def test_is_matching_stem_delimiter_variations(self):
        self.assertTrue(is_matching_stem("2026-09-24__رياض.md", LectureIdentity(date="2026-09-24", keyword="رياض")))
        self.assertTrue(is_matching_stem("2026-09-24  رياض.md", LectureIdentity(date="2026-09-24", keyword="رياض")))

    def test_is_matching_stem_different_keyword_rejects(self):
        self.assertFalse(is_matching_stem("2026-09-19_منطق_AI.md", LectureIdentity(date="2026-09-19", keyword="سيرة")))

    def test_is_matching_stem_different_date_rejects(self):
        self.assertFalse(is_matching_stem("2026-09-20_سيرة.md", LectureIdentity(date="2026-09-19", keyword="سيرة")))

    def test_is_matching_stem_numbered_stem(self):
        self.assertTrue(is_matching_stem("17-فقه_البيوع_17.md", LectureIdentity(keyword="بيوع", stem="17-فقه_البيوع_17")))
        self.assertTrue(is_matching_stem("17-فقه_البيوع_17.md", LectureIdentity(keyword="بيوع", lecture_number="17")))
        self.assertTrue(is_matching_stem("17-فقه_البيوع_17.md", LectureIdentity(date="17", keyword="بيوع")))
        self.assertTrue(is_matching_stem("17-فقه_البيوع_17_AI.md", LectureIdentity(keyword="فقه_البيوع", lecture_number="17")))
        self.assertFalse(is_matching_stem("17-فقه_البيوع_17.md", LectureIdentity(keyword="بيوع", lecture_number="18")))
        self.assertFalse(is_matching_stem("17-فقه_البيوع_17.md", LectureIdentity(keyword="رياض", lecture_number="17")))

    def test_is_matching_stem_numbered_stem_requires_target_num(self):
        # A numbered candidate must NOT match when no target_num or stem is provided
        self.assertFalse(is_matching_stem("17-فقه_البيوع_17.md", LectureIdentity(keyword="بيوع")))
        self.assertFalse(is_matching_stem("17-فقه_البيوع_17.md", LectureIdentity(keyword="فقه البيوع")))


class TestExportNumberedDocuments(unittest.TestCase):
    def test_export_numbered_lecture_documents_e2e(self):
        with tempfile.TemporaryDirectory() as tmp_root, tempfile.TemporaryDirectory() as tmp_onedrive:
            root_path = Path(tmp_root)
            onedrive_base = Path(tmp_onedrive)
            
            raw_dir = root_path / "02_Raw_Inputs"
            matn_dir = root_path / "01_Matn_Sources"
            raw_dir.mkdir(parents=True)
            matn_dir.mkdir(parents=True)
            
            raw_file = raw_dir / "17-فقه_البيوع_17.md"
            matn_file = matn_dir / "17-فقه_البيوع_17.md"
            
            matn_text = "وَيَحْرُمُ رِبَا النَّسِيئَةِ، فِي بَيْعِ كُلِّ جِنْسَيْنِ اتَّفَقَا فِي عِلَّةِ رِبَا الْفَضْلِ."
            matn_file.write_text(matn_text, encoding="utf-8")
            
            transcript = """17-فقه_البيوع_17.mp3
Source guide
المادة: فقه البيوع (سعد الخثلان)
المحاضرة: 17

الحمد لله والصلاة والسلام على رسول الله.
**(ويحرم ربا النسيئة في بيع كل جنسين اتفقا في علة ربا الفضل)**
طالب: صوت غير مسموع.
نعم، هذا ضابط ربا النسيئة."""
            raw_file.write_text(transcript, encoding="utf-8")
            
            results = export_documents(
                transcript,
                original_md_path=raw_file,
                base_onedrive=onedrive_base,
                project_root=root_path,
            )
            
            # 1. Output files exist
            onedrive_file = results["onedrive_file"]
            project_file = results["project_file"]
            self.assertTrue(onedrive_file.exists())
            self.assertTrue(project_file.exists())
            
            # 2. Correct naming
            self.assertEqual(onedrive_file.name, "17-فقه_البيوع_17.docx")
            self.assertEqual(project_file.name, "17-فقه_البيوع_17_AI.md")
            
            # 3. Check docx content and headers
            doc = docx.Document(str(onedrive_file))
            headers = [p.text.strip() for p in doc.paragraphs[:3]]
            self.assertEqual(headers[0], "بسم الله الرحمن الرحيم")
            self.assertEqual(headers[1], "المادة: فقه البيوع (سعد الخثلان)")
            self.assertEqual(headers[2], "المحاضرة: 17")
            
            # 4. Check archiving into processed/
            self.assertFalse(raw_file.exists())
            self.assertFalse(matn_file.exists())
            self.assertTrue((raw_dir / "processed" / "17-فقه_البيوع_17.md").exists())
            self.assertTrue((matn_dir / "processed" / "17-فقه_البيوع_17.md").exists())

    def test_export_documents_metadata_from_path_sets_header_lines(self):
        with tempfile.TemporaryDirectory() as tmp_root, tempfile.TemporaryDirectory() as tmp_onedrive:
            root_path = Path(tmp_root)
            onedrive_base = Path(tmp_onedrive)
            
            raw_dir = root_path / "02_Raw_Inputs"
            raw_dir.mkdir(parents=True)
            raw_file = raw_dir / "17-فقه_البيوع_17.md"
            
            # Markdown text body has NO header block
            transcript = "الحمد لله، نبدأ في شرح ربا النسيئة."
            raw_file.write_text(transcript, encoding="utf-8")
            
            results = export_documents(
                transcript,
                original_md_path=raw_file,
                base_onedrive=onedrive_base,
                project_root=root_path,
            )
            
            onedrive_file = results["onedrive_file"]
            doc = docx.Document(str(onedrive_file))
            headers = [p.text.strip() for p in doc.paragraphs[:3]]
            self.assertEqual(headers[0], "بسم الله الرحمن الرحيم")
            self.assertEqual(headers[1], "المادة: فقه البيوع (سعد الخثلان)")
            self.assertEqual(headers[2], "المحاضرة: 17")

    def test_export_documents_dynamic_stem_fallback_uses_course_subject(self):
        with tempfile.TemporaryDirectory() as tmp_root, tempfile.TemporaryDirectory() as tmp_onedrive:
            root_path = Path(tmp_root)
            onedrive_base = Path(tmp_onedrive)
            
            # Transcript for Diwan al-Shafi'i lecture 5 without mirrored stem in body
            transcript = "المادة: ديوان الشافعي\nالمحاضرة: 5\nقال الإمام الشافعي رحمه الله."
            results = export_documents(
                transcript,
                base_onedrive=onedrive_base,
                project_root=root_path,
            )
            project_file = results["project_file"]
            self.assertIn("5-ديوان_الشافعي_5_AI.md", project_file.name)


class TestExportDocxHelpers(unittest.TestCase):
    def test_extract_candidate_query_strips_outer_parentheses(self):
        self.assertEqual(_extract_candidate_query("(مبادئ كل علم)"), "مبادئ كل علم")
        self.assertEqual(_extract_candidate_query(" (مبادئ كل علم) "), "مبادئ كل علم")
        self.assertEqual(_extract_candidate_query("مبادئ كل علم"), "مبادئ كل علم")
        self.assertEqual(_extract_candidate_query("((متداخل))"), "(متداخل)")

    def test_format_segment_output_primary(self):
        res = MatnMatchResult(
            mode=MatchMode.PRIMARY,
            text="مَبَادِئُ كُلِّ عِلْمٍ",
            start_pos=0,
            end_pos=20,
            is_fallback=False
        )
        self.assertEqual(_format_segment_output(res, is_standalone=True), ["**(مَبَادِئُ كُلِّ عِلْمٍ)**"])
        self.assertEqual(_format_segment_output(res, is_standalone=False), "**(مَبَادِئُ كُلِّ عِلْمٍ)**")

    def test_format_segment_output_sandwiched_fallback(self):
        res = MatnMatchResult(
            mode=MatchMode.FALLBACK,
            text="سقط من المتن",
            start_pos=-1,
            end_pos=-1,
            is_fallback=True,
            is_sandwiched=True
        )
        self.assertEqual(
            _format_segment_output(res, is_standalone=True),
            ["<!-- fallback -->", "**(سقط من المتن)**"]
        )
        self.assertEqual(
            _format_segment_output(res, is_standalone=False),
            "<!-- fallback --> **(سقط من المتن)**"
        )

    def test_format_segment_output_isolated_dialect_fallback(self):
        # Dialect segment without sandwiching fails verification -> unbolded
        res = MatnMatchResult(
            mode=MatchMode.FALLBACK,
            text="ده كلام عامي خالص مش كدا؟",
            start_pos=-1,
            end_pos=-1,
            is_fallback=True,
            is_sandwiched=False
        )
        self.assertEqual(
            _format_segment_output(res, is_standalone=True),
            ["(ده كلام عامي خالص مش كدا؟)"]
        )
        self.assertEqual(
            _format_segment_output(res, is_standalone=False),
            "(ده كلام عامي خالص مش كدا؟)"
        )


class TestLectureIdentity(unittest.TestCase):
    def test_lecture_identity_fields_and_mapping_protocol(self):
        ident = LectureIdentity(
            date="2026-10-09",
            lecture_number="17",
            keyword="بيوع",
            subject_name="فقه البيوع (سعد الخثلان)",
            audio_file="17-فقه_البيوع_17.mp3",
            stem="17-فقه_البيوع_17",
        )
        # Attribute access
        self.assertEqual(ident.date, "2026-10-09")
        self.assertEqual(ident.lecture_number, "17")
        self.assertEqual(ident.keyword, "بيوع")

        # Dict / Mapping access
        self.assertEqual(ident["date"], "2026-10-09")
        self.assertEqual(ident.get("keyword"), "بيوع")
        self.assertEqual(ident.get("missing", "default_val"), "default_val")
        self.assertIn("lecture_number", ident)
        # Methods must not be exposed as dictionary keys
        self.assertNotIn("matches", ident)
        self.assertNotIn("get", ident)
        self.assertEqual(ident.get("matches", "fallback"), "fallback")
        with self.assertRaises(KeyError):
            _ = ident["matches"]

        # Mutability via mapping protocol
        ident["keyword"] = "فقه_البيوع"
        self.assertEqual(ident.keyword, "فقه_البيوع")

    def test_lecture_identity_header_line_3(self):
        # 1. Date precedence
        ident_date = LectureIdentity(date="2026-09-18", lecture_number="5")
        self.assertEqual(ident_date.header_line_3, "2026-09-18")

        # 2. Numbered lecture without date
        ident_num = LectureIdentity(lecture_number="17")
        self.assertEqual(ident_num.header_line_3, "المحاضرة: 17")

        # 3. Fallback when neither is present
        ident_none = LectureIdentity()
        self.assertEqual(ident_none.header_line_3, "تاريخ_غير_محدد")

    def test_lecture_identity_resolved_stem(self):
        # 1. Explicit stem
        ident1 = LectureIdentity(stem="custom-stem_01")
        self.assertEqual(ident1.resolved_stem, "custom-stem_01")

        # 2. Numbered lecture stem derivation
        ident2 = LectureIdentity(lecture_number="17", subject_name="فقه البيوع (سعد الخثلان)")
        self.assertEqual(ident2.resolved_stem, "17-فقه_البيوع_سعد_الخثلان_17")

        # 3. Date fallback
        ident3 = LectureIdentity(date="2026-10-01")
        self.assertEqual(ident3.resolved_stem, "2026-10-01")

    def test_lecture_identity_merge(self):
        base = LectureIdentity(date="2026-10-09")
        other = LectureIdentity(keyword="بيوع", subject_name="فقه البيوع (سعد الخثلان)")
        merged = base.merge(other)
        self.assertEqual(merged.date, "2026-10-09")
        self.assertEqual(merged.keyword, "بيوع")
        self.assertEqual(merged.subject_name, "فقه البيوع (سعد الخثلان)")

    def test_lecture_identity_merge_override_precedence(self):
        # Explicit override in 'other' must supersede non-empty value in 'base'
        base = LectureIdentity(lecture_number="1", subject_name="سيرة قديمة")
        override = LectureIdentity(lecture_number="2", subject_name="سيرة (الرحيق المختوم)")
        merged = base.merge(override)
        self.assertEqual(merged.lecture_number, "2")  # overridden
        self.assertEqual(merged.subject_name, "سيرة (الرحيق المختوم)")  # overridden

        # Overriding a date-based lecture with an explicit numbered lecture clears date to prioritize numbered routing
        base_date = LectureIdentity(date="2026-10-09", subject_name="سيرة قديمة")
        merged_num = base_date.merge(override)
        self.assertEqual(merged_num.lecture_number, "2")
        self.assertIsNone(merged_num.date)

    def test_lecture_identity_matching_date_and_numbered(self):
        # Date match
        ident_date = LectureIdentity(date="2026-09-24", keyword="رياض")
        self.assertTrue(ident_date.matches("2026-09-24__رياض.md"))
        self.assertFalse(ident_date.matches("2026-09-25__رياض.md"))

        # Mirrored numbered match with keyword
        ident_num = LectureIdentity(lecture_number="17", keyword="بيوع")
        self.assertTrue(ident_num.matches("17-فقه_البيوع_17.md"))
        self.assertTrue(ident_num.matches("17-بيوع-17.docx"))
        self.assertFalse(ident_num.matches("18-فقه_البيوع_18.md"))

    def test_lecture_identity_matching_with_subject_name_only(self):
        # Identity instantiated with only subject_name (no keyword) must derive tokens and match stems
        ident_subj = LectureIdentity(lecture_number="17", subject_name="فقه البيوع (سعد الخثلان)")
        self.assertTrue(ident_subj.matches("17-فقه_البيوع_17.md"))
        self.assertTrue(ident_subj.matches("17-بيوع-17.docx"))
        self.assertFalse(ident_subj.matches("18-فقه_البيوع_18.md"))
        self.assertFalse(ident_subj.matches("17-رياض-17.md"))



class TestCanonicalRoutingSSOT(unittest.TestCase):
    def test_canonical_routing_loads_from_context_md(self):
        self.assertGreater(len(CANONICAL_SUBJECT_NAMES), 30)
        # Verify core subjects are mapped dynamically from CONTEXT.md
        self.assertEqual(CANONICAL_SUBJECT_NAMES.get("بيوع"), "فقه البيوع (سعد الخثلان)")
        self.assertEqual(CANONICAL_SUBJECT_NAMES.get("خثلان"), "فقه البيوع (سعد الخثلان)")
        self.assertEqual(CANONICAL_SUBJECT_NAMES.get("فقه_البيوع"), "فقه البيوع (سعد الخثلان)")
        self.assertEqual(CANONICAL_SUBJECT_NAMES.get("سيرة"), "سيرة (الرحيق المختوم)")
        self.assertEqual(CANONICAL_SUBJECT_NAMES.get("رحيق"), "سيرة (الرحيق المختوم)")
        self.assertEqual(CANONICAL_SUBJECT_NAMES.get("منطق"), "المنطق (السلم المنورق)")
        self.assertEqual(CANONICAL_SUBJECT_NAMES.get("سلم"), "المنطق (السلم المنورق)")
        self.assertEqual(CANONICAL_SUBJECT_NAMES.get("زاد"), "زاد المعاد")

    def test_load_canonical_routing_from_custom_table(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            custom_context = Path(tmp_dir) / "CONTEXT.md"
            custom_context.write_text(
                "## 3. توجيه المتون (Routing Map)\n\n"
                "| الكلمة الدلالية | المادة المعتمدة | المتن ومصدره |\n"
                "| --- | --- | --- |\n"
                "| `(عقيدة)` / `(واسطية)` | **العقيدة الواسطية** | الواسطية لشيخ الإسلام |\n",
                encoding="utf-8"
            )
            routing = load_canonical_routing(custom_context)
            self.assertEqual(routing.get("عقيدة"), "العقيدة الواسطية")
            self.assertEqual(routing.get("واسطية"), "العقيدة الواسطية")
            self.assertEqual(routing.get("العقيدة"), "العقيدة الواسطية")
            self.assertEqual(routing.get("الواسطية"), "العقيدة الواسطية")

    def test_load_canonical_routing_with_swapped_columns(self):
        # Table where subject column comes first, keywords second
        with tempfile.TemporaryDirectory() as tmp_dir:
            custom_context = Path(tmp_dir) / "CONTEXT.md"
            custom_context.write_text(
                "## 3. توجيه المتون (Routing Map)\n\n"
                "| المادة المعتمدة | الكلمة الدلالية |\n"
                "| --- | --- |\n"
                "| **أصول الفقه** | `(أصول)` / `(ورقات)` |\n",
                encoding="utf-8"
            )
            routing = load_canonical_routing(custom_context)
            self.assertEqual(routing.get("أصول"), "أصول الفقه")
            self.assertEqual(routing.get("ورقات"), "أصول الفقه")
            self.assertEqual(routing.get("أصول الفقه"), "أصول الفقه")



if __name__ == "__main__":
    unittest.main()
